#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Nyomtatható Word-változat a szakértői kérdőívhez.

Ugyanabból a tételregiszterből és mechanizmus-felosztásból dolgozik, mint az
applikáció (expert_priors.py), így a papír és a webes változat szövege azonos.

    python3 make_expert_form_docx.py   →  PREDICT_szakertoi_prior_urlap.docx
"""

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import os
import sys

from expert_priors import LOWER_CODES, UPPER_CODES, localized_items
from expert_texts import UI, ui_texts

LANG = "en" if "--lang" in sys.argv and sys.argv[sys.argv.index("--lang") + 1] == "en" else "hu"
# --role fogtechnikus: a fogtechnikusi változat (más háttérkérdések, laborból nézett megfogalmazás)
ROLE = "fogtechnikus" if "--role" in sys.argv and sys.argv[sys.argv.index("--role") + 1] == "fogtechnikus" else "fogorvos"
T = ui_texts(LANG, ROLE)
ITEMS = localized_items(LANG)
OUT = {
    ("hu", "fogorvos"): "PREDICT_szakertoi_prior_urlap.docx",
    ("en", "fogorvos"): "PREDICT_expert_prior_form_EN.docx",
    ("hu", "fogtechnikus"): "PREDICT_fogtechnikus_prior_urlap.docx",
    ("en", "fogtechnikus"): "PREDICT_technician_prior_form_EN.docx",
}[(LANG, ROLE)]
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
BLUE = RGBColor(0x1C, 0x5C, 0xAB)
SHADE_HEAD = "DCE8F7"
SHADE_Q = "F2F2F0"
CB = "☐ "
BIZ_LINE = "50 %   ·   60 %   ·   70 %   ·   80 %   ·   90 %   ·   95 %   ·   99 %      " + ("(karikázza be)" if "--lang" not in sys.argv or sys.argv[sys.argv.index("--lang") + 1] != "en" else "(circle one)")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = Cm(1.6)
sec.bottom_margin = Cm(1.5)


def set_font(style, name="Calibri", size=11, bold=None, color=None):
    style.font.name = name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(key), name)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


set_font(doc.styles["Normal"], size=11, color=INK)
doc.styles["Normal"].paragraph_format.space_after = Pt(3)
set_font(doc.styles["Heading 1"], size=15, bold=True, color=BLUE)
set_font(doc.styles["Heading 2"], size=12.5, bold=True, color=INK)
set_font(doc.styles["Heading 3"], size=11, bold=True, color=INK)
for name in ("Heading 1", "Heading 2", "Heading 3"):
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(4)


def add_field(par, instr):
    run = par.add_run()
    run.font.size = Pt(8.5)
    run.font.color.rgb = INK2
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(text)
    run._r.append(end)


header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run("PREDICT · " + (T["login_h1"]) + " · v1.2 (2026-09-06)")
run.font.size = Pt(8.5)
run.font.color.rgb = INK2
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run(("Code" if LANG == "en" else "Kód") + ": ________   ·   ")
run.font.size = Pt(8.5)
run.font.color.rgb = INK2
add_field(footer, "PAGE")
run = footer.add_run(" / ")
run.font.size = Pt(8.5)
run.font.color.rgb = INK2
add_field(footer, "NUMPAGES")


def para(text="", bold=False, italic=False, size=None, color=None, align=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def bullet(text, lead=None):
    p = doc.add_paragraph(style="List Number")
    if lead:
        r = p.add_run(lead)
        r.bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(2)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_text(cell, text, bold=False, size=None, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    for i, line in enumerate(text.split("\n")):
        if i:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
        r = p.add_run(line)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)


def set_widths(table, widths):
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


def two_col(rows, widths=(5.0, 12.4), head=None):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if head:
        r = t.add_row()
        m = r.cells[0].merge(r.cells[1])
        cell_text(m, head, bold=True, size=11.5)
        shade(m, SHADE_HEAD)
    for q, a in rows:
        r = t.add_row()
        cell_text(r.cells[0], q, bold=True, size=10)
        shade(r.cells[0], SHADE_Q)
        cell_text(r.cells[1], a, size=10.5)
    set_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# --- Címlap és útmutató -------------------------------------------------------
HU = LANG == "hu"
if os.path.exists("static/predict-logo.png"):
    # Címlap fejléce, mint a levélben: balra a PREDICT-logó, jobbra a Semmelweis
    # Egyetem logója, keret nélküli kétcellás táblázatban, alulra igazítva.
    if os.path.exists("static/semmelweis-logo.png"):
        header = doc.add_table(rows=1, cols=2)
        header.autofit = False
        left, right = header.rows[0].cells
        left.width, right.width = Cm(8.7), Cm(8.7)
        for cell in (left, right):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
        left.paragraphs[0].add_run().add_picture("static/predict-logo.png", width=Cm(2.6))
        right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right.paragraphs[0].add_run().add_picture("static/semmelweis-logo.png", width=Cm(5.6))
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    else:
        logo_par = doc.add_paragraph()
        logo_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_par.add_run().add_picture("static/predict-logo.png", width=Cm(4.2))
        logo_par.paragraph_format.space_after = Pt(4)
para("PREDICT" + ("-vizsgálat" if HU else " study"), bold=True, size=13, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(T["start_h1"], bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(T["start_eyebrow"], size=13, color=INK2, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
two_col([
    (T["name_label"].replace(" *", ""), "________________________________"),
    (T["affiliation_label"], "________________________________"),
    ("Dátum" if HU else "Date", "2026. ____ . ____ ." if HU else "____ / ____ / 2026"),
    ("Kód (a vizsgálatvezető tölti ki)" if HU else "Code (filled in by the investigator)", "SZ ____"),
], head="Kitöltő" if HU else "Respondent")

doc.add_heading(T["greeting"], level=1)
para(T["intro_1"])
para(T["intro_2_lead"] + " " + T["intro_2"] + " " + ("A kitöltés körülbelül 35–45 perc." if HU else "Completing it takes about 35–45 minutes."), after=6)
doc.add_heading(T["success_h3"], level=2)
para(T["success_def"], bold=True)
para(T["comparable_1"] + " " + T["comparable_lead"] + T["comparable_2"], after=6)
doc.add_heading(T["how_h3"], level=2)
para(T["how_intro"])
for key in ("1", "2", "3", "4"):
    bullet(" " + T[f"how_{key}"], T[f"how_{key}_lead"])
para(("Minden adottságnál leírjuk, hogyan nézzük a vizsgálatban, hogy ugyanarra gondoljunk. A csillaggal jelölt kérdéseket kérjük mindenképpen kitölteni; a többi hasznos, de kihagyható."
      if HU else "For each feature we describe how we record it in the study, so that we are thinking of the same thing. Please answer the questions marked with a star; the others are useful but optional."),
     italic=True, color=INK2, after=6)
doc.add_heading(T["data_h3"], level=2)
para(T["data_p"])
page_break()

# --- Felkészítő (SHELF-mintájú): magyarázat, kidolgozott példa, gyakorlókérdések ---
doc.add_heading(T["prep_h1"], level=1)
para(T["prep_lead"], italic=True, color=INK2)
for key in ("1", "2", "3"):
    doc.add_heading(T[f"prep_{key}_h3"], level=2)
    para(T[f"prep_{key}_p"])
doc.add_heading(T["prep_example_h3"], level=2)
para(T["prep_example_p"], after=6)
doc.add_heading(T["practice_h3"], level=2)
para(T["practice_intro"])
two_col([
    (f"{index}. {kerdes}", f"{T['range_min']} ______   {T['point_label']} ______   {T['range_max']} ______   {egyseg}")
    for index, (_, kerdes, egyseg, _) in enumerate(T["practice_items"], start=1)
] + [(T["practice_dir_q"], "   ".join(CB + label for _, label in T["practice_dir_opts"]) + "\n" + BIZ_LINE)])
para(("A válaszok: " if HU else "The answers: ") + "; ".join(f"{index}. {valasz} {egyseg}" for index, (_, _, egyseg, valasz) in enumerate(T["practice_items"], start=1))
     + ". " + T["practice_dir_answer"] + " " + T["practice_narrow"], size=9, color=INK2, after=6)
page_break()

# --- A. Önről ------------------------------------------------------------------
doc.add_heading(T["a_eyebrow"] + ": " + T["a_h2"].lower() if HU else T["a_eyebrow"] + ": " + T["a_h2"].lower(), level=1)
opts = lambda key: "     ".join(CB + label for _, label in T[key])
two_col(([
    (T["a_kepesites"], "________"),
    (T["a_mester"], opts("a_mester_opts")),
] if ROLE == "fogtechnikus" else [
    (T["a_diploma"], "________"),
    (T["a_szakvizsga"], "________________________________"),
]) + [
    (T["a_evek"], "________"),
    (T["a_fogsorok"], opts("a_fogsorok_opts")),
    (T["a_evi"], "________"),
    (T["a_oktat"], opts("a_oktat_opts")),
    (T["a_tevekenyseg"], opts("a_tevekenyseg_opts") + "  ____________"),
] + ([(T["a_visszajelzes"], opts("a_visszajelzes_opts"))] if ROLE == "fogtechnikus" else []))

# --- B. Általában ----------------------------------------------------------------
doc.add_heading(T["b_h2"], level=1)
two_col([
    ("Sikeres fogsorok aránya *" if HU else "Share of successful dentures *", T["b1"].replace(" *", "") + "\n\n________ " + T["b1_unit"]),
    ("Mennyiben múlik az adottságokon? *" if HU else "How much depends on the anatomy? *", T["b2_help"] + "\n\n________ " + T["b2_unit"]),
    ("Felső vagy alsó?" if HU else "Upper or lower?", T["b3"] + "\n" + opts("b3_opts") + "\n" + T["b3_ratio"] + "  ________"),
    ("Mikor javasolna inkább implantátumon elhorgonyzott fedőlemezes fogpótlást?" if HU else "When would you rather recommend an implant-retained overdenture?", T["b4"] + "\n\n______________________________________________________________\n______________________________________________________________"),
    ("Korábbi fogsor" if HU else "Previous denture", T["b5"] + "\n" + opts("b5_opts") + "   — " + T["b5_ph"] + ": ______________________"),
])
page_break()

# --- C. Adottságok ------------------------------------------------------------------
doc.add_heading(T["items_h2"], level=1)
para(T["items_instruction"], italic=True, color=INK2)
BIZ = "50 %   ·   60 %   ·   70 %   ·   80 %   ·   90 %   ·   95 %   ·   99 %      " + ("(karikázza be)" if HU else "(circle one)")
for index, item in enumerate(ITEMS, start=1):
    head = f"{index}. / {len(ITEMS)}   {item['nev']}   [{item['jaw']} · {item['kod']}]"
    direction = (CB + f"{T['dir_A_small']}:  {item['A']}\n" + CB + f"{T['dir_B_small']}:  {item['B']}\n")
    if item["optimum"]:
        direction += CB + f"{T['dir_opt_small']}: {T['dir_opt']}\n"
    direction += CB + T["dir_none"] + "     " + CB + T["dir_dk"]
    rows = [
        (T["item_how"].rstrip(":"), item["rogzit"]),
        (T["q_direction"], direction),
        (T["q_certainty"], BIZ),
        (T["q_hundred"], T["hundred_help"] + "\n\n" + "\n".join(
            f"{T['variant_' + pole]} ({item[pole]}):   {T['range_min']} ______   {T['point_label']} ______   {T['range_max']} ______   {T['hundred_unit']}"
            for pole in (("A", "M", "B") if "M" in item else ("A", "B")))
         + "\n\n" + T["hundred_help_nodiff"] + "\n" + f"{T['variant_K']}:   {T['range_min']} ______   {T['point_label']} ______   {T['range_max']} ______   {T['hundred_unit']}"
         + "\n\n" + CB + T["magnitude_unknown"] + "     " + T["hundred_hidden_dk"]),
    ]
    if item["kuszob"]:
        rows.append(("Hol a határ?" if HU else "Where is the limit?", item["kuszob"] + "   ______________"))
    for sub in item["subs"]:
        rows.append(("Részletek" if HU else "Details", sub["kerdes"] + "\n" + "   ".join(CB + label for _, label in sub["opciok"])))
    rows.append((T["comment_label"], "______________________________________________________________"))
    two_col(rows, head=head)
    if index % 2 == 0 and index < len(ITEMS):
        page_break()

page_break()
# --- D. Összegzés ----------------------------------------------------------------------
doc.add_heading(T["d_eyebrow"] + ": " + T["d_h2"].lower(), level=1)
codes_upper = "   ".join(f"{item['kod']} = {item['nev']}" for item in ITEMS if item["kod"] in UPPER_CODES)
codes_lower = "   ".join(f"{item['kod']} = {item['nev']}" for item in ITEMS if item["kod"] in LOWER_CODES)
two_col([
    ("Felső állcsont" if HU else "Upper jaw", T["d1_felso"] + " " + T["d1_help"] + "\n\n1. ________   2. ________   3. ________\n\n" + codes_upper),
    ("Alsó állcsont" if HU else "Lower jaw", T["d1_also"] + " " + T["d1_help"] + "\n\n1. ________   2. ________   3. ________\n\n" + codes_lower),
    ("Legrosszabb páros" if HU else "Worst pair", T["d2"] + " " + T["d2_help"] + "\n\n______________________________________________________________"),
    ("Kiegyenlítő adottság" if HU else "Compensating feature", T["d3"] + "\n\n______________________________________________________________"),
    ("Mit hagytunk ki?" if HU else "What did we leave out?", T["d4"] + "\n\n______________________________________________________________"),
    (T["d5"], opts("d5_opts")),
    ("Hozzájárulás" if HU else "Consent", T["consent"] + "\n\n" + ("Dátum" if HU else "Date") + ": ____________________     " + ("Aláírás" if HU else "Signature") + ": ____________________"),
])

zoom = doc.settings.element.find(qn("w:zoom"))
if zoom is not None and zoom.get(qn("w:percent")) is None:
    zoom.set(qn("w:percent"), "100")
doc.save(OUT)
print("saved", OUT)
